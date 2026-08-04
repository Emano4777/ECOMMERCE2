"""Gera apresentação PowerPoint interativa da Licence Farma para licenciados Poupaqui."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets" / "apresentacao_licenciado"
STATIC = ROOT / "static"
OUTPUT = ROOT / "docs" / "Apresentacao_Interativa_Licenciados_Poupaqui.pptx"

RED = RGBColor(205, 7, 48)
DEEP_RED = RGBColor(157, 4, 35)
YELLOW = RGBColor(248, 198, 36)
DARK = RGBColor(23, 25, 35)
TEXT = RGBColor(38, 42, 53)
MUTED = RGBColor(100, 106, 122)
LIGHT = RGBColor(247, 247, 249)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(31, 157, 85)
BLUE = RGBColor(40, 112, 205)

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
links = []


def rect(slide, x, y, w, h, fill, radius=False, line=None, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill; shape.fill.transparency = transparency
    shape.line.color.rgb = line or fill
    if radius:
        try: shape.adjustments[0] = 0.18
        except Exception: pass
    return shape


def text(slide, value, x, y, w, h, size=20, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         font="Aptos", valign=MSO_ANCHOR.MIDDLE, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear(); box.text_frame.word_wrap = True
    box.text_frame.vertical_anchor = valign
    box.text_frame.margin_left = Inches(margin); box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin); box.text_frame.margin_bottom = Inches(margin)
    p = box.text_frame.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value; run.font.name = font; run.font.size = Pt(size)
    run.font.bold = bold; run.font.color.rgb = color
    return box


def add_rich_lines(slide, items, x, y, w, line_h=0.45, size=17):
    for i, (lead, body) in enumerate(items):
        yy = y + i * line_h
        text(slide, lead, x, yy, 0.45, line_h, size=size, color=RED, bold=True, align=PP_ALIGN.CENTER)
        text(slide, body, x + 0.48, yy, w - 0.48, line_h, size=size, color=TEXT)


def add_picture_cover(slide, path, x, y, w, h, dim=0):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    target = w / h; source = iw / ih
    if source > target:
        crop_w = int(ih * target); left = (iw - crop_w) / 2 / iw
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        pic.crop_left = left; pic.crop_right = left
    else:
        crop_h = int(iw / target); top = (ih - crop_h) / 2 / ih
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        pic.crop_top = top; pic.crop_bottom = top
    if dim:
        rect(slide, x, y, w, h, DARK, transparency=dim)
    return pic


def logo(slide, x=0.35, y=0.18, w=1.2):
    p = STATIC / "poupaqui-logo.png"
    if p.exists():
        slide.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(w))
    else:
        text(slide, "poupaqui", x, y, w, 0.4, 18, WHITE, True)


def topbar(slide, title="Licence Farma apresenta"):
    rect(slide, 0, 0, 13.333, 0.68, RED)
    logo(slide, 0.35, 0.12, 1.1)
    text(slide, title, 1.7, 0.11, 7.5, 0.42, 15, WHITE, False)
    text(slide, "Apresentação interativa", 10.55, 0.11, 2.35, 0.42, 12, YELLOW, True, PP_ALIGN.RIGHT)


def bottom_nav(slide, home=True, back_target=None, next_target=None, label="AVANÇAR"):
    if home:
        b = rect(slide, 0.35, 6.87, 1.15, 0.38, DARK, True)
        text(slide, "⌂ MENU", 0.35, 6.87, 1.15, 0.38, 11, WHITE, True, PP_ALIGN.CENTER)
        links.append((b, 1))
    if back_target is not None:
        b = rect(slide, 10.12, 6.87, 1.2, 0.38, WHITE, True, line=RED)
        text(slide, "← VOLTAR", 10.12, 6.87, 1.2, 0.38, 11, RED, True, PP_ALIGN.CENTER)
        links.append((b, back_target))
    if next_target is not None:
        b = rect(slide, 11.47, 6.87, 1.5, 0.38, RED, True)
        text(slide, label, 11.47, 6.87, 1.5, 0.38, 11, WHITE, True, PP_ALIGN.CENTER)
        links.append((b, next_target))


def button(slide, label, x, y, w, h, target=None, external=None, fill=RED, color=WHITE, size=17):
    b = rect(slide, x, y, w, h, fill, True)
    text(slide, label, x, y, w, h, size, color, True, PP_ALIGN.CENTER)
    if target is not None: links.append((b, target))
    if external: b.click_action.hyperlink.address = external
    return b


def pill(slide, label, x, y, w, fill=LIGHT, color=TEXT):
    rect(slide, x, y, w, 0.34, fill, True)
    text(slide, label, x, y, w, 0.34, 11, color, True, PP_ALIGN.CENTER)


def benefit_card(slide, number, title, body, x, y, w=3.7, h=1.35):
    rect(slide, x, y, w, h, WHITE, True, line=RGBColor(232, 233, 238))
    rect(slide, x + 0.18, y + 0.2, 0.48, 0.48, RED, True)
    text(slide, str(number), x + 0.18, y + 0.2, 0.48, 0.48, 16, WHITE, True, PP_ALIGN.CENTER)
    text(slide, title, x + 0.78, y + 0.16, w - 0.95, 0.38, 17, DARK, True)
    text(slide, body, x + 0.78, y + 0.53, w - 0.95, h - 0.64, 12.5, MUTED)


# 1 — CAPA
s = prs.slides.add_slide(blank)
add_picture_cover(s, ASSETS / "home_desktop.png", 0, 0, 13.333, 7.5, dim=32)
rect(s, 0, 0, 6.85, 7.5, DEEP_RED, transparency=10)
logo(s, 0.72, 0.55, 1.55)
pill(s, "LICENCE FARMA → LOJAS LICENCIADAS", 0.75, 1.35, 3.35, fill=YELLOW, color=DARK)
text(s, "Sua loja na jornada\ndigital do consumidor", 0.72, 1.9, 5.5, 1.55, 31, WHITE, True)
text(s, "Clique e experimente como o cliente encontra, escolhe e compra — e veja o valor gerado para a sua unidade.",
     0.75, 3.55, 5.35, 1.1, 18, WHITE)
button(s, "COMEÇAR A EXPERIÊNCIA  →", 0.75, 5.15, 3.55, 0.65, target=1, fill=YELLOW, color=DARK, size=16)
text(s, "Use em modo Apresentação (F5) para ativar os cliques.", 0.75, 6.05, 4.5, 0.4, 11, WHITE)

# 2 — MENU
s = prs.slides.add_slide(blank); topbar(s)
text(s, "Escolha por onde começar", 0.65, 1.0, 6.5, 0.65, 28, DARK, True)
text(s, "Você pode seguir a compra como consumidor ou ir direto aos ganhos da loja.", 0.68, 1.65, 8.6, 0.45, 16, MUTED)
cards = [
    ("1", "Viver a compra", "Navegue pela jornada completa, da busca à confirmação do pedido.", 2, RED),
    ("2", "Entender o ganho", "Veja como cada clique do consumidor cria oportunidade para a loja.", 14, DARK),
    ("3", "Conhecer o mercado", "Dados que mostram por que o canal digital virou parte essencial da farmácia.", 16, BLUE),
]
for i, (num, ttl, body, target, color) in enumerate(cards):
    x = 0.75 + i * 4.15
    b = rect(s, x, 2.45, 3.75, 2.65, WHITE, True, line=RGBColor(226, 228, 234))
    rect(s, x + 0.25, 2.72, 0.62, 0.62, color, True)
    text(s, num, x + 0.25, 2.72, 0.62, 0.62, 20, WHITE, True, PP_ALIGN.CENTER)
    text(s, ttl, x + 0.25, 3.48, 3.2, 0.45, 21, DARK, True)
    text(s, body, x + 0.25, 3.98, 3.15, 0.75, 14, MUTED)
    text(s, "CLIQUE PARA ABRIR  →", x + 0.25, 4.72, 2.5, 0.28, 11, color, True)
    links.append((b, target))
button(s, "Abrir o e-commerce real", 4.93, 5.65, 3.45, 0.58,
       external="https://www.drogariaspoupaqui.com.br", fill=RED, size=15)

# 3 — PERSONA
s = prs.slides.add_slide(blank); topbar(s, "A jornada começa com uma necessidade")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "Imagine uma consumidora perto da sua loja…", 0.7, 1.12, 7.1, 0.65, 28, DARK, True)
text(s, "Mariana precisa repor um produto de cuidado diário. Em vez de sair procurando, ela pega o celular.",
     0.73, 1.88, 6.0, 1.0, 19, TEXT)
add_rich_lines(s, [("✓", "Ela quer rapidez."), ("✓", "Quer saber se está disponível."), ("✓", "Quer escolher entrega ou retirada.")],
               0.75, 3.05, 5.7, 0.58, 18)
rect(s, 7.45, 1.12, 4.7, 4.9, DARK, True)
add_picture_cover(s, ASSETS / "home_mobile.png", 8.67, 1.35, 2.25, 4.4)
text(s, "O próximo clique pode levar\npara a sua loja.", 7.72, 6.03, 4.25, 0.55, 20, RED, True, PP_ALIGN.CENTER)
bottom_nav(s, back_target=1, next_target=3, label="ABRIR O SITE")

# 4 — HOME REAL
s = prs.slides.add_slide(blank)
add_picture_cover(s, ASSETS / "home_desktop.png", 0, 0, 13.333, 7.5)
rect(s, 3.93, 0.20, 4.03, 0.48, WHITE, True, line=YELLOW, transparency=8)
text(s, "CLIQUE NA BUSCA", 4.45, 0.25, 2.95, 0.36, 12, RED, True, PP_ALIGN.CENTER)
b = rect(s, 3.93, 0.15, 4.05, 0.60, WHITE, True, line=YELLOW, transparency=100)
links.append((b, 4))
rect(s, 9.4, 6.35, 3.35, 0.72, DARK, True, transparency=7)
text(s, "A vitrine da rede já está pronta.\nAgora a cliente pesquisa.", 9.58, 6.43, 3.0, 0.52, 13, WHITE, True, PP_ALIGN.CENTER)

# 5 — DIGITANDO BUSCA
s = prs.slides.add_slide(blank)
add_picture_cover(s, ASSETS / "home_desktop.png", 0, 0, 13.333, 7.5, dim=18)
rect(s, 3.55, 1.55, 6.25, 2.55, WHITE, True, line=RGBColor(225, 226, 232))
text(s, "O que você está procurando?", 3.95, 1.87, 5.45, 0.45, 20, DARK, True, PP_ALIGN.CENTER)
rect(s, 4.12, 2.52, 5.1, 0.62, LIGHT, True, line=RGBColor(215, 217, 223))
text(s, "protetor solar", 4.35, 2.52, 3.8, 0.62, 18, TEXT)
text(s, "⌕", 8.38, 2.52, 0.55, 0.62, 24, RED, True, PP_ALIGN.CENTER)
button(s, "VER PRODUTOS PRÓXIMOS  →", 4.32, 3.38, 4.7, 0.55, target=5, size=14)
pill(s, "Busca por texto • voz • foto da embalagem", 4.38, 4.35, 4.6, fill=YELLOW, color=DARK)
bottom_nav(s, back_target=3)

# 6 — CATÁLOGO REAL
s = prs.slides.add_slide(blank)
add_picture_cover(s, ASSETS / "catalogo_desktop.png", 0, 0, 13.333, 7.5)
rect(s, 0.35, 6.42, 5.15, 0.68, DARK, True, transparency=4)
text(s, "Produtos, preços e loja aparecem na mesma tela.", 0.55, 6.48, 4.75, 0.46, 14, WHITE, True)
button(s, "ESCOLHER UMA OFERTA  →", 9.72, 6.4, 3.05, 0.58, target=6, size=14)
b = rect(s, 0.95, 1.34, 2.0, 4.22, WHITE, True, line=YELLOW, transparency=100)
links.append((b, 6))

# 7 — DETALHE DO PRODUTO (JORNADA ILUSTRATIVA)
s = prs.slides.add_slide(blank); topbar(s, "Detalhe do produto")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
pill(s, "JORNADA ILUSTRATIVA", 0.7, 1.0, 1.95, fill=YELLOW, color=DARK)
img = STATIC / "categorias" / "dermocosmeticos.jpg"
rect(s, 0.72, 1.58, 4.45, 4.8, WHITE, True, line=RGBColor(226, 228, 234))
add_picture_cover(s, img, 1.05, 1.92, 3.75, 2.35)
text(s, "Produto de cuidado diário", 5.62, 1.65, 5.9, 0.55, 27, DARK, True)
text(s, "Disponível em uma loja Poupaqui próxima", 5.65, 2.24, 5.65, 0.42, 15, GREEN, True)
text(s, "R$ 39,90", 5.65, 2.93, 2.45, 0.62, 28, RED, True)
text(s, "Preço demonstrativo para simular a jornada.", 5.67, 3.48, 4.1, 0.36, 11, MUTED)
pill(s, "RETIRADA NA LOJA", 5.65, 4.08, 2.12, fill=WHITE, color=DARK)
pill(s, "ENTREGA LOCAL", 7.98, 4.08, 1.9, fill=WHITE, color=DARK)
button(s, "ADICIONAR À CESTA  →", 5.65, 4.82, 3.4, 0.66, target=7, size=16)
text(s, "Para a loja: produto certo + disponibilidade + conveniência = mais chance de conversão.",
     5.65, 5.75, 6.25, 0.56, 15, TEXT, True)
bottom_nav(s, back_target=5)

# 8 — FEEDBACK DE CARRINHO
s = prs.slides.add_slide(blank); topbar(s, "Produto adicionado")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
rect(s, 3.4, 1.45, 6.55, 4.4, WHITE, True, line=RGBColor(226, 228, 234))
rect(s, 5.95, 1.85, 1.45, 1.45, GREEN, True)
text(s, "✓", 5.95, 1.85, 1.45, 1.45, 42, WHITE, True, PP_ALIGN.CENTER)
text(s, "Produto adicionado à cesta!", 4.25, 3.48, 4.85, 0.56, 25, DARK, True, PP_ALIGN.CENTER)
text(s, "A cliente pode continuar comprando ou concluir o pedido.", 4.05, 4.08, 5.25, 0.45, 16, MUTED, False, PP_ALIGN.CENTER)
button(s, "IR PARA A CESTA  →", 5.0, 4.78, 3.3, 0.62, target=8, size=15)
bottom_nav(s, back_target=6)

# 9 — CARRINHO
s = prs.slides.add_slide(blank); topbar(s, "Sua cesta")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "Revise seu pedido", 0.7, 1.05, 5.2, 0.58, 26, DARK, True)
rect(s, 0.72, 1.78, 7.7, 3.65, WHITE, True, line=RGBColor(226, 228, 234))
add_picture_cover(s, img, 1.0, 2.08, 1.6, 1.55)
text(s, "Produto de cuidado diário", 2.92, 2.08, 3.65, 0.42, 18, DARK, True)
text(s, "Vendido por Drogarias Poupaqui", 2.93, 2.55, 3.65, 0.35, 13, MUTED)
pill(s, "−   1   +", 2.93, 3.12, 1.55, fill=LIGHT, color=DARK)
text(s, "R$ 39,90", 6.3, 2.65, 1.55, 0.45, 20, RED, True, PP_ALIGN.RIGHT)
rect(s, 8.77, 1.78, 3.82, 3.65, WHITE, True, line=RGBColor(226, 228, 234))
text(s, "Resumo", 9.08, 2.03, 2.8, 0.42, 19, DARK, True)
text(s, "Produtos", 9.08, 2.75, 1.8, 0.32, 14, MUTED)
text(s, "R$ 39,90", 10.72, 2.75, 1.25, 0.32, 14, DARK, True, PP_ALIGN.RIGHT)
text(s, "Entrega", 9.08, 3.25, 1.8, 0.32, 14, MUTED)
text(s, "A calcular", 10.55, 3.25, 1.42, 0.32, 14, DARK, True, PP_ALIGN.RIGHT)
button(s, "CONTINUAR  →", 9.08, 4.34, 2.88, 0.58, target=9, size=14)
text(s, "A cesta mantém o consumidor conectado à loja até a conclusão.", 0.75, 5.78, 7.5, 0.5, 16, TEXT, True)
bottom_nav(s, back_target=7)

# 10 — IDENTIFICAÇÃO
s = prs.slides.add_slide(blank); topbar(s, "Identificação rápida")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "Entre ou crie sua conta", 0.72, 1.1, 5.0, 0.6, 27, DARK, True)
text(s, "O cadastro conecta compra, acompanhamento e relacionamento futuro.", 0.75, 1.72, 6.2, 0.42, 16, MUTED)
rect(s, 0.75, 2.42, 5.1, 3.45, WHITE, True, line=RGBColor(226, 228, 234))
rect(s, 1.14, 2.82, 4.3, 0.58, LIGHT, True, line=RGBColor(215, 217, 223))
text(s, "seu@email.com", 1.35, 2.82, 3.8, 0.58, 15, MUTED)
rect(s, 1.14, 3.62, 4.3, 0.58, LIGHT, True, line=RGBColor(215, 217, 223))
text(s, "••••••••", 1.35, 3.62, 3.8, 0.58, 17, MUTED)
button(s, "ENTRAR E CONTINUAR  →", 1.14, 4.52, 4.3, 0.62, target=10, size=14)
benefit_card(s, 1, "Histórico", "O cliente encontra pedidos e recompra com facilidade.", 6.55, 2.05, 2.72, 1.45)
benefit_card(s, 2, "Relacionamento", "Cupons e notificações fortalecem a recorrência.", 9.6, 2.05, 2.72, 1.45)
benefit_card(s, 3, "Confiança", "A compra fica vinculada à loja e ao atendimento da rede.", 6.55, 3.85, 2.72, 1.45)
benefit_card(s, 4, "Conveniência", "Dados já salvos tornam a próxima compra mais simples.", 9.6, 3.85, 2.72, 1.45)
bottom_nav(s, back_target=8)

# 11 — ENTREGA OU RETIRADA
s = prs.slides.add_slide(blank); topbar(s, "Como deseja receber?")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "A cliente escolhe o que é mais conveniente", 0.72, 1.08, 7.2, 0.62, 27, DARK, True)
opts = [
    (0.8, "⌂", "Retirar na loja", "Aumenta o fluxo no ponto físico e elimina o custo de entrega."),
    (6.75, "→", "Receber em casa", "Amplia o alcance da unidade para consumidores próximos."),
]
for x, icon, ttl, body in opts:
    b = rect(s, x, 2.05, 5.45, 3.3, WHITE, True, line=RGBColor(222, 224, 230))
    rect(s, x + 0.35, 2.42, 0.85, 0.85, RED, True)
    text(s, icon, x + 0.35, 2.42, 0.85, 0.85, 27, WHITE, True, PP_ALIGN.CENTER)
    text(s, ttl, x + 1.42, 2.37, 3.55, 0.55, 22, DARK, True)
    text(s, body, x + 1.42, 3.02, 3.5, 0.85, 15, MUTED)
    text(s, "SELECIONAR  →", x + 1.42, 4.38, 2.2, 0.36, 12, RED, True)
    links.append((b, 11))
text(s, "A própria loja define raio, valor do frete, pedido mínimo e horários.", 2.15, 5.78, 9.0, 0.5, 17, TEXT, True, PP_ALIGN.CENTER)
bottom_nav(s, back_target=9)

# 12 — PAGAMENTO
s = prs.slides.add_slide(blank); topbar(s, "Pagamento")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "Pagamento simples e integrado", 0.72, 1.05, 6.6, 0.62, 27, DARK, True)
text(s, "A oferta de meios conhecidos reduz o atrito na hora de concluir.", 0.75, 1.7, 6.4, 0.42, 16, MUTED)
methods = [("PIX", "Confirmação rápida", GREEN), ("CARTÃO", "Crédito ou débito", BLUE)]
for i, (ttl, sub, color) in enumerate(methods):
    x = 0.78 + i * 3.35
    rect(s, x, 2.45, 3.0, 1.45, WHITE, True, line=RGBColor(222, 224, 230))
    rect(s, x + 0.25, 2.76, 0.62, 0.62, color, True)
    text(s, "✓", x + 0.25, 2.76, 0.62, 0.62, 19, WHITE, True, PP_ALIGN.CENTER)
    text(s, ttl, x + 1.0, 2.62, 1.65, 0.42, 18, DARK, True)
    text(s, sub, x + 1.0, 3.05, 1.68, 0.32, 12, MUTED)
rect(s, 8.08, 2.05, 4.28, 3.5, WHITE, True, line=RGBColor(222, 224, 230))
text(s, "Resumo do pedido", 8.42, 2.35, 3.35, 0.42, 19, DARK, True)
text(s, "Produto", 8.42, 3.12, 1.8, 0.32, 14, MUTED)
text(s, "R$ 39,90", 10.72, 3.12, 1.1, 0.32, 14, DARK, True, PP_ALIGN.RIGHT)
text(s, "Total demonstrativo", 8.42, 3.65, 2.2, 0.32, 14, MUTED)
text(s, "R$ 39,90", 10.72, 3.65, 1.1, 0.32, 16, RED, True, PP_ALIGN.RIGHT)
button(s, "CONFIRMAR PEDIDO  →", 8.42, 4.5, 3.4, 0.58, target=12, size=14)
text(s, "Jornada ilustrativa: nenhuma cobrança será realizada nesta apresentação.", 0.8, 5.72, 6.6, 0.45, 12, MUTED)
bottom_nav(s, back_target=10)

# 13 — SUCESSO
s = prs.slides.add_slide(blank); topbar(s, "Pedido confirmado")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
rect(s, 2.58, 1.24, 8.18, 4.9, WHITE, True, line=RGBColor(226, 228, 234))
rect(s, 5.92, 1.7, 1.5, 1.5, GREEN, True)
text(s, "✓", 5.92, 1.7, 1.5, 1.5, 44, WHITE, True, PP_ALIGN.CENTER)
text(s, "Pedido confirmado!", 4.25, 3.35, 4.82, 0.58, 28, DARK, True, PP_ALIGN.CENTER)
text(s, "A loja recebe o pedido e a cliente acompanha cada etapa.", 3.8, 4.02, 5.72, 0.45, 16, MUTED, False, PP_ALIGN.CENTER)
pill(s, "RECEBIDO", 3.45, 4.82, 1.5, fill=GREEN, color=WHITE)
pill(s, "EM SEPARAÇÃO", 5.22, 4.82, 1.75, fill=YELLOW, color=DARK)
pill(s, "A CAMINHO", 7.24, 4.82, 1.5, fill=LIGHT, color=MUTED)
pill(s, "ENTREGUE", 9.0, 4.82, 1.28, fill=LIGHT, color=MUTED)
button(s, "VER O QUE A LOJA GANHA  →", 4.75, 5.55, 3.85, 0.52, target=14, size=13)
bottom_nav(s, back_target=11)

# 14 — PWA
s = prs.slides.add_slide(blank); topbar(s, "Poupaqui sempre no celular")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "O site também pode virar um ícone no celular", 0.72, 1.03, 7.2, 0.72, 27, DARK, True)
text(s, "Sem depender da loja de aplicativos: o consumidor instala como PWA e volta com um toque.",
     0.75, 1.75, 6.1, 0.8, 17, TEXT)
add_rich_lines(s, [("1", "Acessa o site pelo navegador."), ("2", "Escolhe “Adicionar à tela inicial”."), ("3", "O ícone Poupaqui fica disponível no celular.")],
               0.75, 2.95, 5.8, 0.66, 17)
rect(s, 8.15, 1.02, 3.55, 5.48, DARK, True)
add_picture_cover(s, ASSETS / "home_mobile.png", 8.78, 1.3, 2.3, 4.85)
text(s, "Mais facilidade para voltar.\nMais chances de recompra.", 0.75, 5.62, 5.95, 0.78, 21, RED, True)
bottom_nav(s, back_target=12, next_target=14, label="GANHOS DA LOJA")

# 15 — GANHO POR ETAPA
s = prs.slides.add_slide(blank); topbar(s, "O que acontece para a loja em cada clique")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "A experiência do consumidor vira oportunidade comercial", 0.7, 1.0, 8.3, 0.62, 27, DARK, True)
steps = [
    ("Busca", "Sua loja passa a ser encontrada."),
    ("Catálogo", "Seu estoque ganha uma vitrine digital."),
    ("Escolha", "Preço e conveniência ajudam a converter."),
    ("Pedido", "A operação recebe uma nova venda."),
    ("Pós-venda", "Histórico e comunicação apoiam a recompra."),
]
for i, (ttl, body) in enumerate(steps):
    x = 0.55 + i * 2.55
    rect(s, x, 2.05, 2.22, 2.25, WHITE, True, line=RGBColor(225, 227, 233))
    rect(s, x + 0.78, 2.34, 0.66, 0.66, RED, True)
    text(s, str(i + 1), x + 0.78, 2.34, 0.66, 0.66, 19, WHITE, True, PP_ALIGN.CENTER)
    text(s, ttl, x + 0.2, 3.15, 1.82, 0.38, 17, DARK, True, PP_ALIGN.CENTER)
    text(s, body, x + 0.18, 3.57, 1.86, 0.55, 12.5, MUTED, False, PP_ALIGN.CENTER)
text(s, "Sua loja física continua sendo o centro da operação. O digital acrescenta uma nova porta de entrada.",
     1.15, 5.05, 11.0, 0.72, 21, RED, True, PP_ALIGN.CENTER)
bottom_nav(s, back_target=13, next_target=15, label="COMO FUNCIONA")

# 16 — RESPONSABILIDADES
s = prs.slides.add_slide(blank); topbar(s, "Licence Farma + loja: parceria para vender")
rect(s, 0, 0.68, 13.333, 6.82, LIGHT)
text(s, "Cada parte cuida do que faz melhor", 0.72, 1.0, 7.0, 0.62, 27, DARK, True)
rect(s, 0.72, 1.82, 5.72, 3.95, WHITE, True, line=RGBColor(225, 227, 233))
text(s, "LICENCE FARMA", 1.08, 2.14, 4.9, 0.45, 20, RED, True)
add_rich_lines(s, [("✓", "Mantém e evolui a plataforma."), ("✓", "Integra catálogo, estoque e serviços."),
                   ("✓", "Apoia implantação, campanhas e suporte."), ("✓", "Gera visão e padrões para a rede.")],
               1.05, 2.82, 4.9, 0.58, 16)
rect(s, 6.9, 1.82, 5.72, 3.95, WHITE, True, line=RGBColor(225, 227, 233))
text(s, "LOJA LICENCIADA", 7.25, 2.14, 4.9, 0.45, 20, DARK, True)
add_rich_lines(s, [("✓", "Mantém preço, estoque e horários corretos."), ("✓", "Recebe, separa e acompanha pedidos."),
                   ("✓", "Cuida da dispensação e do atendimento."), ("✓", "Organiza retirada ou entrega local.")],
               7.22, 2.82, 4.9, 0.58, 16)
text(s, "Tecnologia central + execução local = experiência completa para o consumidor.",
     1.45, 6.03, 10.45, 0.5, 18, RED, True, PP_ALIGN.CENTER)
bottom_nav(s, back_target=14, next_target=16, label="VER O MERCADO")

# 17 — MERCADO
s = prs.slides.add_slide(blank); topbar(s, "Por que entrar agora")
rect(s, 0, 0.68, 13.333, 6.82, DARK)
text(s, "O consumidor farmacêutico já adotou o digital", 0.72, 1.02, 8.2, 0.62, 28, WHITE, True)
rect(s, 0.75, 2.0, 3.75, 2.55, RED, True)
text(s, "R$ 21,58 bi", 1.02, 2.34, 3.2, 0.72, 29, WHITE, True, PP_ALIGN.CENTER)
text(s, "em vendas digitais nas redes acompanhadas", 1.1, 3.15, 3.0, 0.72, 15, WHITE, False, PP_ALIGN.CENTER)
rect(s, 4.8, 2.0, 3.75, 2.55, WHITE, True)
text(s, "+54,82%", 5.07, 2.34, 3.2, 0.72, 29, RED, True, PP_ALIGN.CENTER)
text(s, "de crescimento em 12 meses", 5.2, 3.15, 2.95, 0.72, 15, TEXT, False, PP_ALIGN.CENTER)
rect(s, 8.85, 2.0, 3.75, 2.55, YELLOW, True)
text(s, "+39,85%", 9.12, 2.34, 3.2, 0.72, 29, DARK, True, PP_ALIGN.CENTER)
text(s, "de consumidores aderindo ao canal digital", 9.25, 3.15, 2.95, 0.72, 15, DARK, False, PP_ALIGN.CENTER)
text(s, "Fonte: Abrafarma, dados auditados pela FIA-USP, período dez/2024 a nov/2025.",
     0.78, 4.88, 8.7, 0.4, 11, RGBColor(190, 193, 203))
text(s, "O digital não substitui a loja. Ele amplia onde e quando ela pode ser escolhida.",
     1.12, 5.55, 11.1, 0.72, 21, WHITE, True, PP_ALIGN.CENTER)
bottom_nav(s, back_target=15, next_target=17, label="QUERO PARTICIPAR")

# 18 — CTA
s = prs.slides.add_slide(blank)
add_picture_cover(s, ASSETS / "home_desktop.png", 0, 0, 13.333, 7.5, dim=42)
rect(s, 0, 0, 13.333, 7.5, DEEP_RED, transparency=22)
logo(s, 5.75, 0.65, 1.85)
text(s, "Sua loja está pronta para ganhar\numa nova porta de entrada?", 2.1, 1.85, 9.15, 1.35, 32, WHITE, True, PP_ALIGN.CENTER)
text(s, "A Licence Farma prepara a plataforma. Sua equipe transforma a oportunidade em atendimento e venda.",
     2.3, 3.45, 8.75, 0.8, 18, WHITE, False, PP_ALIGN.CENTER)
button(s, "ABRIR O E-COMMERCE REAL", 3.15, 4.72, 3.45, 0.65,
       external="https://www.drogariaspoupaqui.com.br", fill=YELLOW, color=DARK, size=15)
button(s, "REVER A EXPERIÊNCIA", 6.88, 4.72, 3.2, 0.65, target=2, fill=WHITE, color=RED, size=15)
button(s, "VOLTAR AO MENU", 5.22, 5.65, 2.9, 0.52, target=1, fill=DARK, color=WHITE, size=13)
text(s, "Licence Farma • Poupaqui Ecommerce", 4.65, 6.75, 4.0, 0.35, 12, WHITE, True, PP_ALIGN.CENTER)


# Resolve hyperlinks internos depois que todos os slides existem.
for shape, target in links:
    shape.click_action.target_slide = prs.slides[target]

# Propriedades do arquivo.
prs.core_properties.title = "Poupaqui Ecommerce — Apresentação interativa para licenciados"
prs.core_properties.subject = "Jornada do consumidor e proposta de valor da Licence Farma para lojas licenciadas"
prs.core_properties.author = "Licence Farma"
prs.core_properties.keywords = "Poupaqui, ecommerce, licenciado, jornada do consumidor, apresentação interativa"
prs.core_properties.comments = "Execute em modo Apresentação (F5) para usar os botões e a navegação interna."
prs.save(OUTPUT)
print(OUTPUT)
print(f"slides={len(prs.slides)} links={len(links)}")
