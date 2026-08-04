"""Gera a planilha de custos Poupaqui em XLSX e CSV compatível com Excel pt-BR."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
XLSX = DOCS / "Custos_Ecommerce_Poupaqui.xlsx"
CSV = DOCS / "Custos_Ecommerce_Poupaqui.csv"

CAMBIO = 5.0638

ROWS = [
    ["Infraestrutura", "Vercel", "Plano Pro", "Mensal", 20.00, 101.28, "Sim", "Licence Farma", "Inclui crédito de uso; excedentes variáveis"],
    ["Dados", "Supabase", "Plano Pro recomendado", "Mensal", 25.00, 126.60, "Sim para produção", "Licence Farma", "Confirmar o plano atualmente contratado"],
    ["Mensageria", "WasenderAPI", "Basic — 1 sessão", "Mensal", 6.00, 30.38, "Condicional", "Licence Farma", "Confirmar o número de sessões necessárias"],
    ["Inteligência artificial", "Anthropic", "Reserva inicial", "Mensal variável", 10.00, 50.64, "Condicional", "Licence Farma", "Medir tokens por recurso"],
    ["Pesquisa", "Serper", "50 mil créditos / 6 meses", "Compra avulsa", 50.00, 253.19, "Condicional", "Licence Farma", "Equivale a cerca de US$ 8,33/mês se distribuído em seis meses"],
    ["Imagens", "Cloudinary", "Free", "Franquia", 0.00, 0.00, "Condicional", "Licence Farma", "25 créditos mensais"],
    ["Imagens", "Cloudinary", "Plus", "Mensal", 99.00, 501.32, "Opcional", "Licence Farma", "Upgrade conforme armazenamento, tráfego e transformações"],
    ["E-mail", "Resend", "Free", "Franquia", 0.00, 0.00, "Condicional", "Licence Farma", "3.000 e-mails/mês e limite de 100/dia"],
    ["E-mail", "Resend", "Pro", "Mensal", 20.00, 101.28, "Opcional", "Licence Farma", "50 mil e-mails mensais"],
    ["OCR", "OCR.space", "Free", "Franquia", 0.00, 0.00, "Condicional", "Licence Farma", "25 mil solicitações/mês; 500/dia por IP"],
    ["OCR", "OCR.space", "Pro", "Mensal", 30.00, 151.91, "Opcional", "Licence Farma", "300 mil solicitações mensais"],
    ["Mapas", "Google Maps", "APIs", "Por uso", None, None, "Condicional", "Licence Farma", "Configurar orçamento, alertas e limites"],
    ["Pagamento", "Mercado Pago", "Taxa por transação", "Por venda", None, None, "Sim, se habilitado", "Loja", "Varia conforme meio de pagamento e prazo de recebimento"],
    ["Marketplace", "Mercado Livre", "Comissão e frete", "Por venda", None, None, "Não", "Loja", "Varia conforme categoria e tipo de anúncio"],
    ["Domínio", "Registro .com.br", "Renovação", "Anual", None, None, "Sim", "Licence Farma", "Preencher o valor real da fatura"],
    ["Operação", "Entrega", "Motoboy ou parceiro", "Por entrega ou mensal", None, None, "Condicional", "Loja", "Definir custo, raio e repasse"],
    ["Operação", "Suporte", "Equipe humana", "Mensal", None, None, "Sim", "Licence Farma e loja", "Separar atendimento de níveis 1, 2 e 3"],
    ["Marketing", "Mídia e criação", "Campanhas", "Mensal variável", None, None, "Não", "Licence Farma e loja", "Definir orçamento por região"],
]

HEADERS = ["Categoria", "Fornecedor", "Item", "Cobrança", "Referência (US$)", "Referência (R$)", "Obrigatório", "Responsável sugerido", "Observação"]
RED = "C8102E"
YELLOW = "F5C842"
DARK = "171923"
LIGHT = "F7F7F9"
WHITE = "FFFFFF"
GREEN = "DFF2E1"


def style_title(ws, title, subtitle, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=RED)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    c = ws.cell(2, 1, subtitle)
    c.font = Font(name="Aptos", size=10, italic=True, color="555555")
    c.fill = PatternFill("solid", fgColor="FFF4CF")
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 28


def configure_table(ws, header_row, end_col, end_row):
    thin = Side(style="thin", color="D9D9DF")
    for cell in ws[header_row]:
        if cell.column <= end_col:
            cell.font = Font(name="Aptos", bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=DARK)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row + 1, max_row=end_row, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.fill = PatternFill("solid", fgColor=WHITE if cell.row % 2 else LIGHT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{end_row}"
    ws.freeze_panes = f"A{header_row + 1}"


def make_costs(wb):
    ws = wb.active
    ws.title = "Custos"
    style_title(ws, "Custos do Poupaqui Ecommerce", "Valores de referência em 24/07/2026. Confirme os planos e faturas antes de decidir.", len(HEADERS))
    for col, value in enumerate(HEADERS, 1):
        ws.cell(4, col, value)
    for row_no, row in enumerate(ROWS, 5):
        for col_no, value in enumerate(row, 1):
            ws.cell(row_no, col_no, value)
    configure_table(ws, 4, len(HEADERS), 4 + len(ROWS))
    for row in range(5, 5 + len(ROWS)):
        ws.cell(row, 5).number_format = 'US$ #,##0.00'
        ws.cell(row, 6).number_format = 'R$ #,##0.00'
    widths = [22, 20, 27, 21, 17, 17, 19, 24, 62]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.sheet_view.showGridLines = False


def make_summary(wb):
    ws = wb.create_sheet("Resumo", 0)
    style_title(ws, "Resumo financeiro", "Cenário profissional enxuto e cenários de crescimento.", 6)
    ws["A4"] = "Premissa"; ws["B4"] = "Valor"
    ws["A5"] = "Câmbio de referência"; ws["B5"] = CAMBIO; ws["B5"].number_format = 'R$ 0.0000'
    ws["A6"] = "Data-base"; ws["B6"] = "24/07/2026"
    ws["A8"] = "Operação profissional enxuta"
    ws["A8"].font = Font(size=14, bold=True, color=RED)
    items = [
        ("Vercel Pro", 20.00), ("Supabase Pro", 25.00), ("WasenderAPI Basic", 6.00),
        ("Reserva inicial de IA", 10.00), ("Serper amortizado em 6 meses", 8.33),
    ]
    ws.append([])
    row = 9
    ws.cell(row, 1, "Item"); ws.cell(row, 2, "US$/mês"); ws.cell(row, 3, "R$/mês")
    for name, usd in items:
        row += 1
        ws.cell(row, 1, name); ws.cell(row, 2, usd); ws.cell(row, 3, f"=B{row}*$B$5")
    row += 1
    ws.cell(row, 1, "Subtotal"); ws.cell(row, 2, f"=SUM(B10:B{row-1})"); ws.cell(row, 3, f"=SUM(C10:C{row-1})")
    for c in ws[row]:
        c.font = Font(bold=True, color=DARK); c.fill = PatternFill("solid", fgColor=YELLOW)
    for r in range(10, row + 1):
        ws.cell(r, 2).number_format = 'US$ #,##0.00'; ws.cell(r, 3).number_format = 'R$ #,##0.00'

    start = row + 3
    ws.cell(start, 1, "Cenário de IA"); ws.cell(start, 2, "Mínimo (US$)"); ws.cell(start, 3, "Máximo (US$)"); ws.cell(start, 4, "Mínimo (R$)"); ws.cell(start, 5, "Máximo (R$)")
    scenarios = [("Piloto", 7, 12), ("Crescimento", 32, 48), ("Escala", 130, 180)]
    for i, (name, low, high) in enumerate(scenarios, start + 1):
        ws.cell(i, 1, name); ws.cell(i, 2, low); ws.cell(i, 3, high)
        ws.cell(i, 4, f"=B{i}*$B$5"); ws.cell(i, 5, f"=C{i}*$B$5")
        for col in range(2, 6): ws.cell(i, col).number_format = 'R$ #,##0.00' if col >= 4 else 'US$ #,##0.00'
    configure_table(ws, start, 5, start + len(scenarios))
    ws.column_dimensions["A"].width = 36
    for col in "BCDE": ws.column_dimensions[col].width = 18
    ws.column_dimensions["F"].width = 4
    ws.sheet_view.showGridLines = False


def make_ai(wb):
    ws = wb.create_sheet("Estimativa IA")
    style_title(ws, "Estimativa de inteligência artificial", "Edite a quantidade mensal para simular o custo. Valores são estimativas técnicas.", 8)
    headers = ["Recurso", "Chamadas/mês", "Tokens entrada", "Tokens saída", "US$/M entrada", "US$/M saída", "Custo US$", "Custo R$"]
    for col, value in enumerate(headers, 1): ws.cell(4, col, value)
    data = [
        ("Busca inteligente", 2000, 700, 400, 1, 5),
        ("Suporte com IA", 300, 1000, 500, 1, 5),
        ("Cross-sell", 100, 2500, 300, 1, 5),
        ("Personalização curta", 500, 300, 80, 1, 5),
        ("Descrição de produto", 500, 500, 320, 1, 5),
        ("Visão de receita/embalagem", 50, 3000, 200, 3, 15),
    ]
    for r, values in enumerate(data, 5):
        for c, value in enumerate(values, 1): ws.cell(r, c, value)
        ws.cell(r, 7, f"=B{r}*((C{r}/1000000)*E{r}+(D{r}/1000000)*F{r})")
        ws.cell(r, 8, f"=G{r}*Resumo!$B$5")
        ws.cell(r, 7).number_format = 'US$ #,##0.00'; ws.cell(r, 8).number_format = 'R$ #,##0.00'
    total = 5 + len(data)
    ws.cell(total, 1, "TOTAL"); ws.cell(total, 7, f"=SUM(G5:G{total-1})"); ws.cell(total, 8, f"=SUM(H5:H{total-1})")
    for c in range(1, 9):
        ws.cell(total, c).font = Font(bold=True); ws.cell(total, c).fill = PatternFill("solid", fgColor=YELLOW)
    configure_table(ws, 4, 8, total)
    widths = [32, 17, 18, 16, 17, 17, 16, 16]
    for idx, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(idx)].width = width
    ws.sheet_view.showGridLines = False


def write_csv():
    # UTF-8 com BOM + ponto e vírgula: abre corretamente no Excel em português.
    with CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(HEADERS)
        for row in ROWS:
            writer.writerow(["" if value is None else str(value).replace(".", ",") if isinstance(value, float) else value for value in row])


def main():
    wb = Workbook()
    make_costs(wb)
    make_summary(wb)
    make_ai(wb)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(XLSX)
    write_csv()
    print(XLSX)
    print(CSV)


if __name__ == "__main__":
    main()
