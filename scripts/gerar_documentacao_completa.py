"""Gera a versão Word da documentação mestre Poupaqui a partir do Markdown."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Documentacao_Completa_Ecommerce_Poupaqui.md"
OUTPUT = ROOT / "docs" / "Documentacao_Completa_Ecommerce_Poupaqui.docx"
RED = "C8102E"
YELLOW = "F5C842"


def clean_inline(text):
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, rows):
    parsed = [[clean_inline(x.strip()) for x in row.strip().strip("|").split("|")] for row in rows]
    if len(parsed) > 1 and all(re.fullmatch(r":?-{3,}:?", x.replace(" ", "")) for x in parsed[1]):
        parsed.pop(1)
    cols = max(len(row) for row in parsed)
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(parsed):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    if r_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if r_idx == 0:
                shade(cell, RED)
            elif r_idx % 2 == 0:
                shade(cell, "F5F5F5")
    doc.add_paragraph()


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = ROOT / "static" / "poupaqui-logo.png"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Inches(1.3))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("POUPAQUI ECOMMERCE")
    r.bold = True; r.font.size = Pt(30); r.font.color.rgb = RGBColor.from_string(RED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Documento mestre de operação, custos, implantação e proposta de valor")
    r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor(45, 45, 55)
    doc.add_paragraph()
    p = doc.add_paragraph("Versão 1.0 • 24 de julho de 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("www.drogariaspoupaqui.com.br")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7); section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(9.5)
    for name, size in (("Title", 28), ("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11)):
        styles[name].font.name = "Arial"; styles[name].font.size = Pt(size)
        styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor.from_string(RED)

    add_cover(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---":
            i += 1; continue
        if line.startswith("# "):
            i += 1; continue  # título já está na capa
        if line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=1); i += 1; continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=2); i += 1; continue
        if line.startswith("#### "):
            doc.add_heading(clean_inline(line[5:]), level=3); i += 1; continue
        image = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line.strip())
        if image:
            path = SOURCE.parent / image.group(2)
            if path.exists():
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Inches(6.5))
                cap = doc.add_paragraph(image.group(1)); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cap.runs:
                    run.italic = True; run.font.size = Pt(8); run.font.color.rgb = RGBColor(100, 100, 100)
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            add_table(doc, rows); continue
        if line.startswith("> "):
            p = doc.add_paragraph(clean_inline(line[2:]))
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(8)
            for run in p.runs:
                run.italic = True; run.font.color.rgb = RGBColor(80, 80, 80)
            i += 1; continue
        if re.match(r"^- ", line):
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet"); i += 1; continue
        m = re.match(r"^\d+\.\s+(.*)", line)
        if m:
            doc.add_paragraph(clean_inline(m.group(1)), style="List Number"); i += 1; continue
        p = doc.add_paragraph(clean_inline(line))
        p.paragraph_format.space_after = Pt(5)
        i += 1

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Poupaqui Ecommerce • Documento estratégico e operacional")
        run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(RED)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
